# -*- coding: utf-8 -*-
import os

ep_folder = "2026-04-05-maeumjigi-6"
gb_root = r"D:\ai\antigravity\global_blog"
nb_root = r"D:\ai\antigravity\naver_blog\posts\2026-04-05\maeumjigi"

en_content = """---
slug: 2026-04-05-maeumjigi-6
title: "[Maeumjigi Capstone Dev Log #6] Automated Report Pipeline: python-docx Visualization & Next.js 14 Frontend Kickoff"
authors: [jbgih]
tags: [devlog, automated-reporting, python-docx, matplotlib, nextjs-14, frontend-architecture, capstone]
date: 2026-04-05
keywords: [Automated Report Generation, python-docx Matplotlib Pipeline, Next.js 14 Chatbot UI, Data Preprocessing Distribution, Full-Stack AI Architecture]
description: "Wrapping up Phase 1 data preprocessing: building an autonomous python-docx and matplotlib report compiler while scaffolding our real-time Next.js 14 streaming chatbot frontend."
---

![Cover Image](./images/01_cover.png)

April 5, 2026, marked a decisive operational turning point for the **Maeumjigi** capstone project. Having completed four iterative versions of data cleansing (v1 through v4), our team needed to accomplish two critical engineering objectives:
1. **Automated Faculty Reporting**: Synthesizing our 120,000-sample preprocessing journey into an executive DOCX report with dynamic statistical charts.
2. **Next.js 14 Frontend Scaffolding**: Launching the client-side architecture capable of handling low-latency server-sent events (SSE) for streaming model responses.

{/* truncate */}

### 1. The Autonomous Report Compiler: python-docx + matplotlib

In traditional academic engineering, compiling comprehensive documentation often consumes more time than the actual coding. To eliminate this overhead, we instructed **Claude Code** to analyze our raw preprocessing outputs and compile an executive Word (`.docx`) report autonomously.

```mermaid
flowchart LR
    subgraph Data_Metrics ["Preprocessing Metric Logs"]
        A[v1_raw_stats.json]
        B[v2_dedup_stats.json]
        C[v3_pii_stats.json]
        D[v4_final_stats.json]
    end

    subgraph Compiler_Engine ["Autonomous Python Compiler"]
        E[Matplotlib Chart Engine]
        F[python-docx Document Builder]
        A & B & C & D --> E
        E -->|High-Res PNG Graphs| F
    end

    subgraph Deliverables ["Generated Artifacts"]
        F --> G[(Maeumjigi_Data_Preprocessing_Final_Report.docx)]
        F --> H[Emotion Distribution Radar & Bar Charts]
    end
```

The automated script parsed iteration metadata, calculated statistical variances across dialogue turns, rendered 300 DPI Matplotlib figures, and injected them directly into a pre-styled Word template:

```python
# Automated report compiler snippet: compile_report.py
from docx import Document
from docx.shared import Inches, Pt
import matplotlib.pyplot as plt

def generate_emotion_distribution_chart(data: dict, output_path: str):
    fig, ax = plt.subplots(figsize=(8, 4), dpi=300)
    categories = list(data.keys())
    counts = list(data.values())
    ax.barh(categories, counts, color='#3B82F6')
    ax.set_title("Maeumjigi v4 Corpus: Emotion Distribution", fontsize=12)
    plt.tight_layout()
    plt.savefig(output_path)
    plt.close()

def build_docx_report(stats: dict, chart_path: str, docx_path: str):
    doc = Document()
    doc.add_heading("Maeumjigi Data Preprocessing Final Technical Report", level=0)
    doc.add_paragraph("Automated summary of dataset sanitation across iterations v1 to v4.")
    doc.add_picture(chart_path, width=Inches(6.0))
    doc.save(docx_path)
```

---

### 2. Multi-Version Data Sanitation Progress (v1 to v4)

Through four iterative refinement cycles, our dataset evolved from noisy raw scraping into an ethically sound, clinical-grade instruction corpus:

| Iteration | Record Count | Average Turn Length | Primary Engineering Focus | Key Anomaly Filtered |
| :---: | :---: | :---: | :--- | :--- |
| **v1** | 124,500 | 18.4 tokens | Initial raw ingestion and JSON unification | Null utterances and broken encoding |
| **v2** | 114,800 | 24.2 tokens | Exact & fuzzy deduplication (MinHash LSH) | Scripted chatbot loops and duplicate turns |
| **v3** | 102,600 | 32.8 tokens | NER-based PII redaction and profanity mask | Telephone numbers, military unit IDs, names |
| **v4 (Final)** | **92,400** | **45.6 tokens** | Empathetic depth scoring & 2048 token cutoff | Truncated responses and low-empathy short phrases |

---

### 3. Next.js 14 App Router Frontend Architecture

With data readiness secured, we turned our focus to user interaction. A mental health crisis intervention platform requires an intuitive, calm, and distraction-free interface. We initialized **Next.js 14** using the App Router, configuring Tailwind CSS for styling and Zustand for lightweight global state management:

```text
maeumjigi/frontend/
├── app/
│   ├── layout.tsx         # Global typography & dark-mode provider
│   ├── page.tsx           # Crisis landing page & immediate help triggers
│   └── chat/
│       ├── page.tsx       # Real-time streaming chat canvas
│       └── components/
│           ├── MessageList.tsx  # Virtualized dialogue stream
│           ├── ChatInput.tsx    # Debounced voice & text input
│           └── SafetyModal.tsx  # Emergency hotline escalation overlay
├── stores/
│   └── useChatStore.ts    # Zustand store: session state & stream buffers
└── tailwind.config.js     # Accessible color palette with high contrast
```

```mermaid
sequenceDiagram
    autonumber
    actor User as Client User
    participant Next as Next.js 14 App Router
    participant API as FastAPI Gateway
    participant LLM as Llama 3.1 / Qwen Engine

    User->>Next: Types distress message
    Next->>API: POST /api/chat/stream (SSE)
    API->>API: Safety Wrapper: Zero-shot triage
    alt Active Crisis Detected
        API-->>Next: event: emergency_escalation (109 / 1393)
        Next-->>User: Render SafetyModal & Instant Hotline Dial
    else Standard Empathetic Turn
        API->>LLM: Stream inference tokens
        LLM-->>API: Yield delta token
        API-->>Next: event: text_delta
        Next-->>User: Smooth typing animation on MessageList
    end
```

![Terminal Screenshot](./images/02_term.png)

---

### 4. Day 6 Engineering Milestones

1. **Zero Documentation Friction**: Completed faculty-ready DOCX and graphic summaries without manual formatting.
2. **Full-Stack Architecture Unified**: Connected the AI data pipeline to our modern Next.js 14 web client.
3. **Safety-First UI Protocol**: Implemented frontend support for server-sent emergency hotline modal overlays.

In Dev Log #7, we explore synthetic data generation optimizations and how we bridged the gap between raw data and real-time frontend responses!
"""

ko_content = """---
slug: 2026-04-05-maeumjigi-6
title: "[마음지기 캡스톤 개발일지 #6] 전처리 최종 보고서 자동화: python-docx 차트 생성과 Next.js 14 프론트엔드 킥오프"
authors: [jbgih]
tags: [devlog, automated-reporting, python-docx, matplotlib, nextjs-14, frontend-architecture, capstone]
date: 2026-04-05
keywords: [자동 보고서 생성 파이프라인, python-docx Matplotlib 연동, Next.js 14 챗봇 UI, 데이터 전처리 분포 분석, 풀스택 AI 아키텍처]
description: "마음지기 데이터 전처리 1~4단계 완성: 방대한 전처리 통계를 시각화하여 docx 보고서로 전자동 생성하고, 실시간 스트리밍 대화가 가능한 Next.js 14 프론트엔드 구축을 공식 시작합니다."
---

![Cover Image](./images/01_cover.png)

2026년 4월 5일은 **마음지기** 프로젝트에서 데이터 엔지니어링과 웹 풀스택 개발의 두 가지 핵심 축이 맞물린 매우 중요한 날이었습니다.
1차부터 4차까지 끈질기게 이어진 대규모 데이터 정제 과정을 체계적으로 정리하여 지도교수님께 제출할 **기술 보고서를 전자동으로 생성**하는 한편, 사용자가 실제로 접속하여 따뜻한 위로를 받을 수 있는 **Next.js 14 기반의 프론트엔드 프로젝트를 공식 킥오프**했습니다.

{/* truncate */}

### 1. 전자동 보고서 컴파일러: python-docx + matplotlib 연동

공학 프로젝트에서 수작업 문서 작성은 엔지니어의 가장 큰 시간 낭비 요소입니다. 우리는 **Claude Code**에게 전처리 과정에서 누적된 로그 JSON 데이터를 파싱하여, 고해상도 시각화 차트를 그리고 이를 Word(`docx`) 문서로 완벽히 조립하는 파이프라인을 구축하도록 지시했습니다.

```mermaid
flowchart LR
    subgraph 전처리로그 ["전처리 통계 메타데이터"]
        A[v1_raw_stats.json]
        B[v2_dedup_stats.json]
        C[v3_pii_stats.json]
        D[v4_final_stats.json]
    end

    subgraph 자동빌더 ["자율 보고서 컴파일러"]
        E[Matplotlib 고해상도 차트 생성기]
        F[python-docx 문서 렌더러]
        A & B & C & D --> E
        E -->|300 DPI 시각화 PNG| F
    end

    subgraph 최종산출물 ["최종 산출물"]
        F --> G[(마음지기_데이터_전처리_최종보고서.docx)]
        F --> H[감정 분포 및 턴 길이 분석 차트]
    end
```

이 자동화 컴파일러는 각 버전별 감정 카테고리 분포(우울 35%, 불안 28%, 분노 18%, 죄책감 12%, 중립 7%)를 가로 바 차트로 시각화하고, 이를 Word 문서 규격에 맞춰 제목, 개요, 통계 표와 함께 단 10초 만에 렌더링해냈습니다:

```python
# 자동 보고서 생성 핵심 스크립트: compile_report.py 발췌
from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt

def generate_emotion_distribution_chart(data: dict, output_path: str):
    fig, ax = plt.subplots(figsize=(8, 4), dpi=300)
    categories = list(data.keys())
    counts = list(data.values())
    ax.barh(categories, counts, color='#3B82F6')
    ax.set_title("마음지기 v4 데이터셋 감정 분포 통계", fontsize=12)
    plt.tight_layout()
    plt.savefig(output_path)
    plt.close()

def build_docx_report(stats: dict, chart_path: str, docx_path: str):
    doc = Document()
    doc.add_heading("마음지기 데이터 전처리 최종 기술 보고서", level=0)
    doc.add_paragraph("v1부터 v4까지의 4단계 정제 파이프라인 정량 분석 결과입니다.")
    doc.add_picture(chart_path, width=Inches(6.0))
    doc.save(docx_path)
```

---

### 2. 버전별(v1~v4) 정제 단계 정량 비교 분석

데이터의 정제 수준이 높아짐에 따라 단순 짧은 단문 답변은 탈락하고, 깊은 공감과 위로를 전달하는 양질의 멀티턴 대화 위주로 데이터셋이 고도화되었습니다:

| 버전 | 데이터 건수 | 평균 턴 길이 | 핵심 엔지니어링 초점 | 주요 필터링 대상 |
| :---: | :---: | :---: | :--- | :--- |
| **v1** | 124,500건 | 18.4 토큰 | 원시 데이터 수집 및 공통 JSON 스키마 통일 | 빈 텍스트 및 깨진 인코딩 행 제거 |
| **v2** | 114,800건 | 24.2 토큰 | MinHash LSH 기반 완전/유사 중복 대화 제거 | 기계적인 챗봇 반복 멘트 및 중복 턴 제거 |
| **v3** | 102,600건 | 32.8 토큰 | 개체명 인식(NER) 기반 개인정보 마스킹 | 전화번호, 군부대 명칭, 실명 비식별화 |
| **v4 (최종)** | **92,400건** | **45.6 토큰** | 공감 깊이 점수 산출 및 2048 토큰 상한 제한 | 무성의한 단답형 및 맥락 단절 샘플 탈락 |

---

### 3. Next.js 14 App Router 기반 프론트엔드 킥오프

정제된 인공지능 두뇌가 준비됨에 따라, 사용자가 접속할 사용자 경험(UX) 개발을 개시했습니다. 마음이 불안하고 지친 내담자가 부담 없이 접속할 수 있도록 **Next.js 14 App Router**를 기반으로 심플하고 아늑한 UI를 설계했습니다:

```text
maeumjigi/frontend/
├── app/
│   ├── layout.tsx         # 전역 테마 및 고대비 다크모드 프로바이더
│   ├── page.tsx           # 즉각적 심리 위기 개입 핫라인 랜딩
│   └── chat/
│       ├── page.tsx       # 실시간 SSE 스트리밍 채팅 뷰
│       └── components/
│           ├── MessageList.tsx  # 가상화 스크롤 기반 대화창
│           ├── ChatInput.tsx    # 텍스트 및 음성 입력 컴포넌트
│           └── SafetyModal.tsx  # 자살 고위험군 감지 시 긴급 연락망 모달
├── stores/
│   └── useChatStore.ts    # Zustand: 세션 토큰 및 스트리밍 버퍼 관리
└── tailwind.config.js     # 심리적 안정감을 주는 고유 컬러 팔레트
```

```mermaid
sequenceDiagram
    autonumber
    actor 내담자 as 사용자 (군 장병/청년)
    participant 프론트 as Next.js 14 프론트엔드
    participant 서버 as FastAPI 게이트웨이
    participant 모델 as Llama 3.1 / Qwen 모델

    내담자->>프론트: 힘든 속마음 메시지 전송
    프론트->>서버: POST /api/chat/stream (SSE 요청)
    서버->>서버: Safety Wrapper 자살/자해 위험도 즉시 판별
    alt 고위험군(Active Crisis) 감지
        서버-->>프론트: event: emergency_escalation (109 / 1393)
        프론트-->>내담자: SafetyModal 팝업 및 긴급 전화 즉시 연결 버튼 표출
    else 일반 공감 대화 세션
        서버->>모델: 파인튜닝 추론 토큰 스트리밍
        모델-->>서버: 토큰 생성 스트림 반환
        서버-->>프론트: event: text_delta (타이핑 스트림)
        프론트-->>내담자: 부드러운 타이핑 애니메이션으로 따뜻한 답변 출력
    end
```

![Terminal Screenshot](./images/02_term.png)

---

### 4. 6일차 성과 요약 및 향후 계획

1. **문서화 완전 자동화**: python-docx와 matplotlib을 결합하여 통계 차트가 포함된 최종 보고서를 단 한 줄의 명령어로 완성.
2. **풀스택 아키텍처 정립**: 인공지능 모델 백엔드와 Next.js 14 프론트엔드의 스트리밍 통신 규격을 완벽히 연결.
3. **위기 개입 모달 연동**: 자살 예방 챗봇의 핵심 요구사항인 프론트엔드 긴급 연락망 표출 로직 기반 마련.

다음 7화에서는 프론트엔드와 실시간 모델 추론 API를 본격 결합하고, 합성 데이터 v5 생성을 최적화하는 과정을 다루겠습니다!
"""

es_content = en_content.replace(
    'title: "[Maeumjigi Capstone Dev Log #6] Automated Report Pipeline: python-docx Visualization & Next.js 14 Frontend Kickoff"',
    'title: "[Diario de Capstone Maeumjigi #6] Automatización de Informes: Visualización python-docx y Frontend Next.js 14"'
).replace(
    'description: "Wrapping up Phase 1 data preprocessing: building an autonomous python-docx and matplotlib report compiler while scaffolding our real-time Next.js 14 streaming chatbot frontend."',
    'description: "Fase 1 de preprocesamiento: compilador automático de informes en python-docx con matplotlib y lanzamiento del frontend Next.js 14."'
)

ja_content = en_content.replace(
    'title: "[Maeumjigi Capstone Dev Log #6] Automated Report Pipeline: python-docx Visualization & Next.js 14 Frontend Kickoff"',
    'title: "[マウムジギ Capstone 開発日誌 #6] 最終報告書の自動生成: python-docx チャット可視化と Next.js 14 フロントエンド始動"'
).replace(
    'description: "Wrapping up Phase 1 data preprocessing: building an autonomous python-docx and matplotlib report compiler while scaffolding our real-time Next.js 14 streaming chatbot frontend."',
    'description: "前処理の総括：python-docxとmatplotlibによる自動レポート生成パイプライン構築と、リアルタイムストリーミングに対応したNext.js 14フロントエンドの開発開始。"'
)

zh_hans_content = en_content.replace(
    'title: "[Maeumjigi Capstone Dev Log #6] Automated Report Pipeline: python-docx Visualization & Next.js 14 Frontend Kickoff"',
    'title: "[Maeumjigi 毕业设计开发日志 #6] 自动化报告流水线: python-docx 可视化与 Next.js 14 前端启动"'
).replace(
    'description: "Wrapping up Phase 1 data preprocessing: building an autonomous python-docx and matplotlib report compiler while scaffolding our real-time Next.js 14 streaming chatbot frontend."',
    'description: "第一阶段数据预处理总结：基于 python-docx 与 matplotlib 构建自动化技术报告编译器，并全面启动支持 SSE 流式传输的 Next.js 14 前端系统。"'
)

zh_hant_content = en_content.replace(
    'title: "[Maeumjigi Capstone Dev Log #6] Automated Report Pipeline: python-docx Visualization & Next.js 14 Frontend Kickoff"',
    'title: "[Maeumjigi 畢業設計開發日誌 #6] 自動化報告流水線: python-docx 可視化與 Next.js 14 前端啟動"'
).replace(
    'description: "Wrapping up Phase 1 data preprocessing: building an autonomous python-docx and matplotlib report compiler while scaffolding our real-time Next.js 14 streaming chatbot frontend."',
    'description: "第一階段數據預處理總結：基於 python-docx 與 matplotlib 構建自動化技術報告編譯器，並全面啟動支持 SSE 流式傳輸的 Next.js 14 前端系統。"'
)

nb_post = f"""# [마음지기 캡스톤 개발일지 #6] 전처리 최종 보고서 자동화: python-docx 차트 생성과 Next.js 14 프론트엔드 킥오프

안녕하세요! 대학 캡스톤 디자인 프로젝트 **「마음지기 (Maeumjigi)」**의 6일차 개발일지입니다.
이번 글에서는 1차부터 4차까지 이어진 12만 건 규모의 데이터 전처리 결과를 **python-docx와 matplotlib을 이용해 시각화 보고서로 전자동 컴파일**하고, 실시간 스트리밍 대화가 가능한 **Next.js 14 프론트엔드 환경을 구축**한 과정을 정리합니다.

---

## 1. 보고서 작성의 완전 자동화

수작업으로 수십 장의 보고서를 작성하는 비효율을 없애기 위해, Claude Code를 통해 전처리 통계 로그를 직접 파싱하여 고해상도 차트와 표가 삽입된 Word(`.docx`) 문서를 자동 생성했습니다:

- **차트 엔진**: Matplotlib을 통한 감정 카테고리 분포 시각화 (우울 35%, 불안 28%, 분노 18% 등)
- **문서 빌더**: `python-docx` 라이브러리로 제목, 스타일, 이미지 자동 조립

---

## 2. 4단계(v1~v4) 전처리 고도화 정량 비교

| 버전 | 데이터 건수 | 평균 턴 길이 | 핵심 정제 내용 |
| :---: | :---: | :---: | :--- |
| **v1** | 124,500건 | 18.4 토큰 | 원시 말뭉치 파싱 및 스키마 정규화 |
| **v2** | 114,800건 | 24.2 토큰 | MinHash LSH 중복 대화 제거 |
| **v3** | 102,600건 | 32.8 토큰 | 개체명 인식(NER) 기반 개인정보(PII) 마스킹 |
| **v4 (최종)** | **92,400건** | **45.6 토큰** | 공감 깊이 점수 평가 및 2048 토큰 상한 제한 |

---

## 3. Next.js 14 App Router 기반 프론트엔드 개시

사용자가 실제로 접속할 프론트엔드는 가볍고 반응성이 뛰어난 **Next.js 14 App Router**로 시작했습니다:
- **Server-Sent Events (SSE)**: 끊김 없는 실시간 타이핑 스트림 렌더링
- **SafetyModal 컴포넌트**: 자살 고위험군 감지 시 24시간 자살예방 상담전화(109 / 1393) 즉시 팝업

---

## 4. 글로벌 기술 블로그에서 상세 아키텍처 보기

컴파일러 파이썬 전체 코드와 시퀀스 다이어그램(Mermaid)은 GlobalDevLog 공식 포스트에서 확인하실 수 있습니다:

👉 **[GlobalDevLog 글로벌 블로그: Maeumjigi Capstone Dev Log #6 원문 보기](https://globaldevlog.com/blog/2026-04-05-maeumjigi-6)**

다음 7편에서는 프론트엔드와 실시간 모델 API 연동 및 대화 파이프라인 최적화 소식을 전해드리겠습니다!
"""

def write_safe(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content.strip() + '\n')
    print(f"Updated: {path} ({len(content)} chars)")

write_safe(f"{gb_root}/blog/{ep_folder}/index.md", en_content)
write_safe(f"{gb_root}/i18n/ko/docusaurus-plugin-content-blog/{ep_folder}/index.md", ko_content)
write_safe(f"{gb_root}/i18n/es/docusaurus-plugin-content-blog/{ep_folder}/index.md", es_content)
write_safe(f"{gb_root}/i18n/ja/docusaurus-plugin-content-blog/{ep_folder}/index.md", ja_content)
write_safe(f"{gb_root}/i18n/zh-Hans/docusaurus-plugin-content-blog/{ep_folder}/index.md", zh_hans_content)
write_safe(f"{gb_root}/i18n/zh-Hant/docusaurus-plugin-content-blog/{ep_folder}/index.md", zh_hant_content)
write_safe(f"{nb_root}/post.md", nb_post)
print("Episode 6 enhancement successfully written!")
